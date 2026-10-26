# app/documents.py
# Handles PDF and DOCX text extraction and splitting into chunks.
# Uses PyMuPDF for PDFs and python-docx for DOCX files.

import math
from pathlib import Path

import pymupdf  # PyMuPDF (formerly imported as 'fitz')
from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter

from app.config import settings


# Allowed file extensions
ALLOWED_EXTENSIONS = {".pdf", ".docx"}


def validate_file_extension(filename: str) -> str:
    """
    Check the file extension is supported.
    Returns the lowercase extension or raises ValueError.
    """
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type '{ext}'. Supported types: {', '.join(ALLOWED_EXTENSIONS)}"
        )
    return ext


def extract_text_from_pdf(file_path: Path) -> str:
    """
    Extract readable text from a PDF using PyMuPDF.
    Returns the combined text from all pages.
    """
    text_parts: list[str] = []
    with pymupdf.open(str(file_path)) as doc:
        for page in doc:
            page_text = page.get_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts)


def extract_text_from_docx(file_path: Path) -> str:
    """
    Extract readable text from a DOCX file using python-docx.
    Returns the combined text from all paragraphs.
    """
    doc = DocxDocument(str(file_path))
    text_parts = [para.text for para in doc.paragraphs if para.text.strip()]
    return "\n".join(text_parts)


def extract_text(file_path: Path, extension: str) -> str:
    """
    Route to the correct extractor based on file extension.
    Raises ValueError if no text is found.
    """
    if extension == ".pdf":
        text = extract_text_from_pdf(file_path)
    elif extension == ".docx":
        text = extract_text_from_docx(file_path)
    else:
        raise ValueError(f"No extractor for extension '{extension}'.")

    if not text.strip():
        raise ValueError(
            "No readable text could be extracted from this document. "
            "Scanned PDFs may need OCR, which is not supported in this version."
        )
    return text


def count_pages(file_path: Path, extension: str) -> int | None:
    """
    Count pages (PDF) or non-empty paragraphs (DOCX).
    Returns None if counting is not supported for the file type.
    """
    try:
        if extension == ".pdf":
            with pymupdf.open(str(file_path)) as doc:
                return len(doc)
        elif extension == ".docx":
            doc = DocxDocument(str(file_path))
            return len([p for p in doc.paragraphs if p.text.strip()])
    except Exception:
        return None
    return None


def split_text_into_chunks(
    text: str,
    source_filename: str,
    grade: str | None = None,
    subject: str | None = None,
) -> list[dict]:
    """
    Split text into overlapping chunks for retrieval.
    Each chunk is a dict with 'content' and 'metadata'.

    The RecursiveCharacterTextSplitter tries to split on paragraph/sentence
    boundaries first, keeping chunks semantically coherent.
    """
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.CHUNK_SIZE,
        chunk_overlap=settings.CHUNK_OVERLAP,
        length_function=len,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    raw_chunks = splitter.split_text(text)

    chunks = []
    total = len(raw_chunks)
    half = math.ceil(total / 2.0) if total > 0 else 0

    for i, chunk_text in enumerate(raw_chunks):
        meta = {
            "source": source_filename,
            "chunk_index": i,
            "total_chunks": total,
            "is_first_half": (i < half),
        }
        if grade:
            meta["grade"] = grade
        if subject:
            meta["subject"] = subject

        chunks.append({
            "content": chunk_text,
            "metadata": meta,
        })
    return chunks
